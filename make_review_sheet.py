"""핫딜 리뷰 기록표(xlsx) 만들기. deals_checked.json의 제품·최저가 판매처로 칸을 미리 채운다.

사람이 판매처 페이지에서 최근 리뷰 50건을 읽고(2026-10-23 결정) 노란 칸을 채운 뒤,
reviews_from_sheet.py로 reviews_checked.json을 만든다. 이미 있는 기록표는 덮어쓰지 않는다.
"""
import datetime
import json
import pathlib
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = pathlib.Path(__file__).parent
READ_N = 50  # 제품마다 읽을 최근 리뷰 수(2026-10-23 결정)
# 평가 요소: [이름, 좋았다는 말 예, 아쉬웠다는 말 예]. app.template.html의 ASPECTS와 이름이 같아야 한다
ASPECTS = {
    '기저귀': [('샘', '안 새요, 밤새 괜찮아요', '샜어요, 등·옆으로 샘'),
            ('피부', '발진 없어요, 순해요', '발진, 빨개짐, 짓무름'),
            ('사이즈·밴드', '잘 맞아요', '작아요·커요, 허리·허벅지 자국'),
            ('흡수', '보송해요, 흡수 잘돼요', '축축해요, 무거워져요'),
            ('두께·통기', '얇아요, 시원해요', '두꺼워요, 더워요, 땀띠')],
    '물티슈': [('두께', '도톰해요', '얇아요, 잘 찢어져요'),
            ('물기', '촉촉해요', '물이 적어요·말라요, 너무 흥건해요'),
            ('성분·향', '향 없어요, 순해요', '향이 나요, 따가워요, 트러블'),
            ('뽑기·캡', '한 장씩 잘 나와요', '줄줄이 딸려 나와요, 캡이 안 닫혀요'),
            ('보풀·잔여감', '보풀 없어요', '보풀, 미끈거려요')],
}
FIXED = ['번호', '제품명', '최저가 판매처', '가격 확인', '판매처에서 열기',
         '리뷰 확인 날짜', '판매처 평점', '판매처 리뷰 수', '읽은 리뷰 수']
FONT = 'Arial'
F_HEAD = PatternFill('solid', fgColor='1B2333')
F_SUB = PatternFill('solid', fgColor='F4F6FA')
F_IN = PatternFill('solid', fgColor='FFF4D1')  # 채우는 칸
THIN = Side(style='thin', color='E3E6EE')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def font(**kw):
    return Font(name=FONT, size=kw.pop('size', 10), **kw)


def guide(ws, checked_text):
    ws.title = '읽는 법'
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 34
    ws.column_dimensions['C'].width = 40
    rows = [
        ('애보사 핫딜 리뷰 기록표', None, None, 'title'),
        (f"제품과 최저가 판매처는 '지금 가격'({checked_text} 확인) 기준으로 미리 채웠습니다.", None, None, 'note'),
        ('', None, None, None),
        ('채우는 칸', '연한 노란 칸만 채웁니다', '회색 글자 칸(제품명·판매처·가격 확인·링크)은 고치지 않습니다', 'row'),
        ('1. 열기', "'판매처에서 열기'를 눌러 그 판매처 상품 페이지를 엽니다", '판매처가 기록표와 다르면 그 줄은 비워 둡니다(앱에서 숨겨짐)', 'row'),
        ('2. 평점·리뷰 수', '판매처가 보여 주는 평점과 리뷰 수를 그대로 적습니다', '평점은 0~5, 소수점 그대로(예: 4.8)', 'row'),
        ('3. 최근 리뷰 읽기', f'리뷰를 최신순으로 놓고 {READ_N}건을 읽습니다', f"리뷰가 {READ_N}건보다 적으면 있는 만큼 읽고 '읽은 리뷰 수'에 실제 건수를 적습니다", 'row'),
        ('4. 요소별 세기', '요소마다 좋았다는 말과 아쉬웠다는 말을 셉니다', '리뷰 한 건이 여러 요소를 말하면 요소마다 1씩, 한 요소에서 좋은 말·아쉬운 말이 둘 다 나오면 양쪽에 1씩', 'row'),
        ('5. 날짜', "다 읽은 날을 '리뷰 확인 날짜'에 적습니다", '예: 2026-10-24', 'row'),
        ('6. 확인 칸', "맨 오른쪽 '확인'이 '완료'가 되면 다 채운 것입니다", "'빈칸 있음'·'건수 확인'이 뜨면 그 줄을 다시 봅니다", 'row'),
        ('', None, None, None),
        ('세지 않는 것', '판매자 답글, 같은 사람이 반복해 쓴 글, 사진만 있고 글이 없는 리뷰', '이런 리뷰는 50건에 넣지 않고 다음 리뷰를 읽습니다', 'row'),
        ('', None, None, None),
        ('평가 요소', '좋았다는 말로 세는 것(예)', '아쉬웠다는 말로 세는 것(예)', 'head'),
    ]
    for kind, items in ASPECTS.items():
        rows.append((kind, None, None, 'sub'))
        rows += [(k, g, b, 'row') for k, g, b in items]
    rows += [('', None, None, None),
             ('예시 줄(가짜 값)', '아래는 칸을 어떻게 채우는지 보여 주는 예시입니다', '실제 기록표 시트에는 이런 줄이 없습니다. 이 시트의 값은 앱에 쓰지 않습니다', 'note2'),
             ('리뷰 확인 날짜', '2026-10-24', None, 'ex'),
             ('판매처 평점 · 리뷰 수', '4.8 · 1234', None, 'ex'),
             ('읽은 리뷰 수', '50', None, 'ex'),
             ('샘', '좋았다 21 · 아쉬웠다 4', None, 'ex')]
    for i, (a, b, c, kind) in enumerate(rows, start=1):
        ws.cell(i, 1, a); ws.cell(i, 2, b); ws.cell(i, 3, c)
        for col in range(1, 4):
            cell = ws.cell(i, col)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
            cell.font = font()
        if kind == 'title':
            ws.cell(i, 1).font = font(size=14, bold=True)
        elif kind in ('note', 'note2'):
            ws.cell(i, 1).font = font(bold=kind == 'note2')
            ws.cell(i, 2).font = font(color='5D6678')
            ws.cell(i, 3).font = font(color='5D6678')
        elif kind == 'head':
            for col in range(1, 4):
                ws.cell(i, col).font = font(bold=True, color='FFFFFF')
                ws.cell(i, col).fill = F_HEAD
        elif kind == 'sub':
            ws.cell(i, 1).font = font(bold=True)
            for col in range(1, 4):
                ws.cell(i, col).fill = F_SUB
        elif kind == 'row':
            ws.cell(i, 1).font = font(bold=True)
            for col in range(1, 4):
                ws.cell(i, col).border = BOX
        elif kind == 'ex':
            ws.cell(i, 2).fill = F_IN
            ws.cell(i, 2).font = font(color='5D6678', italic=True)
    ws.merge_cells('A1:C1')
    ws.merge_cells('A2:C2')


def sheet(wb, kind, items, checked_text):
    ws = wb.create_sheet(kind)
    asp = ASPECTS[kind]
    n_fixed = len(FIXED)
    first_asp = n_fixed + 1
    last_asp = n_fixed + 2 * len(asp)
    c_check, c_memo = last_asp + 1, last_asp + 2
    # 1행: 묶음 제목, 2행: 칸 제목
    for j, h in enumerate(FIXED, start=1):
        ws.cell(1, j, h)
        ws.merge_cells(start_row=1, start_column=j, end_row=2, end_column=j)
    for a, (name, good, bad) in enumerate(asp):
        col = first_asp + 2 * a
        ws.cell(1, col, name)
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 1)
        ws.cell(2, col, '좋았다')
        ws.cell(2, col + 1, '아쉬웠다')
        ws.cell(2, col).comment = Comment(f'좋았다는 말(예): {good}', '애보사')
        ws.cell(2, col + 1).comment = Comment(f'아쉬웠다는 말(예): {bad}', '애보사')
    for col, h in ((c_check, '확인'), (c_memo, '메모')):
        ws.cell(1, col, h)
        ws.merge_cells(start_row=1, start_column=col, end_row=2, end_column=col)
    for r in (1, 2):
        for col in range(1, c_memo + 1):
            cell = ws.cell(r, col)
            cell.font = font(bold=True, color='FFFFFF')
            cell.fill = F_HEAD
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = BOX
    ws.cell(3, 1, f"노란 칸만 채웁니다. 제품·판매처는 '지금 가격'({checked_text} 확인) 기준. 최근 리뷰 {READ_N}건을 읽고 요소별 건수를 적습니다.")
    ws.cell(3, 1).font = font(color='5D6678', italic=True)
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=c_memo)
    first = 4
    for i, d in enumerate(items):
        r = first + i
        vals = [i + 1, d['title'], d['mall'], checked_text]
        for j, v in enumerate(vals, start=1):
            ws.cell(r, j, v).font = font(color='5D6678')
        link = ws.cell(r, 5, '열기')
        link.hyperlink = d['link']
        link.font = font(color='2A6FD6', underline='single')
        for col in range(6, last_asp + 1):
            ws.cell(r, col).fill = F_IN
            ws.cell(r, col).font = font()
        ws.cell(r, 6).number_format = 'yyyy-mm-dd'
        ws.cell(r, 7).number_format = '0.0#'
        ws.cell(r, 8).number_format = '#,##0'
        fa, la = get_column_letter(first_asp), get_column_letter(last_asp)
        ws.cell(r, c_check, f'=IF(COUNTA(F{r}:I{r})+COUNT({fa}{r}:{la}{r})=0,"",'
                            f'IF(COUNTA(F{r}:I{r})+COUNT({fa}{r}:{la}{r})<{4 + 2 * len(asp)},"빈칸 있음",'
                            f'IF(MAX({fa}{r}:{la}{r})>I{r},"건수 확인","완료")))')
        ws.cell(r, c_check).font = font(bold=True)
        ws.cell(r, c_memo).font = font()
        for col in range(1, c_memo + 1):
            ws.cell(r, col).border = BOX
            ws.cell(r, col).alignment = Alignment(vertical='center', wrap_text=col in (2, c_memo))
    last = first + len(items) - 1
    rng = lambda c: f'{get_column_letter(c)}{first}:{get_column_letter(c)}{last}'
    dvs = [
        (DataValidation(type='date', operator='between', formula1='DATE(2026,1,1)', formula2='DATE(2030,12,31)',
                        error='날짜를 2026-10-24처럼 적어 주세요', showErrorMessage=True), [rng(6)]),
        (DataValidation(type='decimal', operator='between', formula1='0', formula2='5',
                        error='평점은 0~5 사이 숫자입니다', showErrorMessage=True), [rng(7)]),
        (DataValidation(type='whole', operator='greaterThanOrEqual', formula1='0',
                        error='리뷰 수는 0 이상 정수입니다', showErrorMessage=True), [rng(8)]),
        (DataValidation(type='whole', operator='between', formula1='1', formula2=str(READ_N),
                        error=f'읽은 리뷰 수는 1~{READ_N}입니다', showErrorMessage=True), [rng(9)]),
        (DataValidation(type='whole', operator='between', formula1='0', formula2=str(READ_N),
                        error=f'건수는 0~{READ_N} 사이 정수입니다', showErrorMessage=True),
         [f'{get_column_letter(first_asp)}{first}:{get_column_letter(last_asp)}{last}']),
    ]
    for dv, ranges in dvs:
        for x in ranges:
            dv.add(x)
        ws.add_data_validation(dv)
    widths = {1: 6, 2: 38, 3: 12, 4: 15, 5: 9, 6: 13, 7: 10, 8: 12, 9: 11, c_check: 11, c_memo: 24}
    for col in range(1, c_memo + 1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(col, 9)
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 20
    ws.freeze_panes = ws.cell(first, 3)


def main():
    deals = json.loads((HERE / 'deals_checked.json').read_text(encoding='utf-8'))
    out = HERE / f'리뷰_기록표_{datetime.date.today():%Y-%m-%d}.xlsx'
    if out.exists() and '--force' not in sys.argv:
        sys.exit(f'{out.name}이 이미 있습니다. 채운 값을 지키려고 덮어쓰지 않습니다(새로 만들려면 --force).')
    wb = Workbook()
    guide(wb.active, deals['checkedAtText'])
    for g in deals['groups']:
        sheet(wb, g['kind'], g.get('items') or g.get('rows'), deals['checkedAtText'])
    wb.save(out)
    print(out.name, ' · '.join(f"{g['kind']} {len(g.get('items') or g.get('rows'))}개" for g in deals['groups']))


if __name__ == '__main__':
    main()
