"""채운 리뷰 기록표(xlsx) → reviews_checked.json. 다 채운 줄만 넣고, 문제가 있으면 줄 번호와 함께 알리고 멈춘다.

사용: python3 reviews_from_sheet.py 리뷰_기록표_2026-10-23.xlsx [--out 다른경로.json(시험용)]
- 칸 위치·요소 이름은 make_review_sheet.py와 같다.
- 지금 deals_checked.json의 제품·최저가 판매처와 다른 줄은 넣되 경고한다(앱에서는 판매처가 같을 때만 보인다).
"""
import datetime
import json
import pathlib
import sys

from openpyxl import load_workbook

from make_review_sheet import ASPECTS, FIXED, READ_N

HERE = pathlib.Path(__file__).parent
FIRST_ROW = 4


def as_date(v):
    if isinstance(v, datetime.datetime):
        return v.date()
    if isinstance(v, datetime.date):
        return v
    try:
        return datetime.date.fromisoformat(str(v).strip()[:10])
    except ValueError:
        return None


def main():
    if len(sys.argv) < 2:
        sys.exit('사용: python3 reviews_from_sheet.py 리뷰_기록표_YYYY-MM-DD.xlsx')
    wb = load_workbook(sys.argv[1], data_only=True)
    deals = json.loads((HERE / 'deals_checked.json').read_text(encoding='utf-8'))
    now = {(d['title'], d['mall']) for g in deals['groups'] for d in (g.get('items') or g.get('rows'))}
    items, errors, warns, empty = [], [], [], 0
    for kind, asp in ASPECTS.items():
        if kind not in wb.sheetnames:
            errors.append(f'{kind} 시트가 없습니다')
            continue
        ws = wb[kind]
        n = len(FIXED)
        for r in range(FIRST_ROW, ws.max_row + 1):
            title, mall = ws.cell(r, 2).value, ws.cell(r, 3).value
            if not title:
                continue
            day, rating, count, read = (ws.cell(r, c).value for c in (6, 7, 8, 9))
            nums = [ws.cell(r, n + 1 + j).value for j in range(2 * len(asp))]
            filled = [v not in (None, '') for v in (day, rating, count, read, *nums)]
            where = f'{kind} {r}행({title})'
            if not any(filled):
                empty += 1
                continue
            if not all(filled):
                errors.append(f'{where}: 빈칸이 있습니다')
                continue
            d = as_date(day)
            try:
                rating, count, read = float(rating), int(count), int(read)
                nums = [int(v) for v in nums]
            except (TypeError, ValueError):
                errors.append(f'{where}: 숫자 칸에 숫자가 아닌 값이 있습니다')
                continue
            bad = []
            if d is None:
                bad.append('날짜 형식')
            if not 0 <= rating <= 5:
                bad.append('평점은 0~5')
            if count < 0:
                bad.append('리뷰 수는 0 이상')
            if not 1 <= read <= READ_N:
                bad.append(f'읽은 리뷰 수는 1~{READ_N}')
            if any(v < 0 or v > read for v in nums):
                bad.append('요소별 건수는 0~읽은 리뷰 수')
            if bad:
                errors.append(f'{where}: ' + ', '.join(bad))
                continue
            if (title, mall) not in now:
                warns.append(f'{where}: 지금 가격의 최저가 판매처와 달라 앱에서는 보이지 않습니다')
            items.append({'title': title, 'mall': mall, 'kind': kind, 'checkedAt': d.isoformat(),
                          'checkedAtText': f'{d.month}월 {d.day}일', 'rating': round(rating, 2), 'ratingCount': count, 'read': read,
                          'aspects': {name: [nums[2 * j], nums[2 * j + 1]] for j, (name, _, _) in enumerate(asp)}})
    for w in warns:
        print('경고 —', w)
    if errors:
        sys.exit('기록표를 고쳐 주세요(reviews_checked.json은 바꾸지 않았습니다):\n- ' + '\n- '.join(errors))
    out = {'source': '최저가 판매처 상품평(사람이 직접 읽고 셈)', 'readPerItem': READ_N, 'items': items}
    dest = pathlib.Path(sys.argv[sys.argv.index('--out') + 1]) if '--out' in sys.argv else HERE / 'reviews_checked.json'
    dest.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding='utf-8')
    print(f'{dest.name}: {len(items)}개 제품 · 비어 있는 줄 {empty}개')


if __name__ == '__main__':
    main()
